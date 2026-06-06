"""Hisobot eksporti: Excel (openpyxl) va PDF (weasyprint).

Og'ir kutubxonalar (openpyxl/weasyprint) funksiya ichida import qilinadi — shunda
bot ishga tushishi ularning tizim bog'liqliklariga bog'liq bo'lmaydi.
"""
import io
import logging
from datetime import date
from decimal import Decimal

from app.db import acquire
from app.handlers.common import fmt_money

log = logging.getLogger("imed-finance-bot.export")


# --------------------------- ma'lumot yig'ish ---------------------------
async def _collect() -> dict:
    """Hisobotlar uchun kerakli barcha kesimlarni bazadan yig'adi."""
    async with acquire() as conn:
        transactions = await conn.fetch(
            """SELECT t.occurred_at, t.kind, t.amount, t.currency, t.amount_uzs,
                      w.name AS wallet, f.name AS fund,
                      COALESCE(s.name, c.name) AS category, t.description
               FROM transactions t
               JOIN wallets w ON w.id = t.wallet_id
               LEFT JOIN funds f ON f.id = t.fund_id
               LEFT JOIN income_sources s ON s.id = t.source_id
               LEFT JOIN expense_categories c ON c.id = t.category_id
               ORDER BY t.occurred_at DESC"""
        )
        pnl = await conn.fetch(
            """SELECT to_char(date_trunc('month', occurred_at), 'YYYY-MM') AS month,
                      COALESCE(SUM(amount_uzs) FILTER (WHERE kind = 'income'), 0) AS income,
                      COALESCE(SUM(amount_uzs) FILTER (WHERE kind = 'expense'), 0) AS expense
               FROM transactions
               WHERE kind IN ('income', 'expense')
               GROUP BY 1 ORDER BY 1 DESC LIMIT 12"""
        )
        wallets = await conn.fetch(
            "SELECT name, currency, balance_uzs FROM v_wallet_balances ORDER BY name"
        )
        funds = await conn.fetch("SELECT name, balance_uzs FROM v_fund_balances ORDER BY name")
        by_category = await conn.fetch(
            """SELECT c.name, COALESCE(SUM(t.amount_uzs), 0) AS total
               FROM transactions t JOIN expense_categories c ON c.id = t.category_id
               WHERE t.kind = 'expense'
               GROUP BY c.name ORDER BY total DESC"""
        )
        debts = await conn.fetch(
            """SELECT direction, counterparty_name, currency, remaining, due_date, status
               FROM v_debt_outstanding WHERE remaining > 0 ORDER BY due_date NULLS LAST"""
        )
    return {
        "transactions": transactions,
        "pnl": pnl,
        "wallets": wallets,
        "funds": funds,
        "by_category": by_category,
        "debts": debts,
    }


# --------------------------- Excel ---------------------------
async def build_excel() -> bytes:
    """Excel hisobotini (bytes) yasaydi: Tranzaksiyalar, P&L, Balanslar, Qarzlar."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    data = await _collect()
    wb = Workbook()

    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Tranzaksiyalar"
    headers = ["Sana", "Tur", "Summa", "Valyuta", "UZS", "Kassa", "Fond", "Kategoriya", "Izoh"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = bold
    for t in data["transactions"]:
        ws.append([
            t["occurred_at"].strftime("%Y-%m-%d %H:%M") if t["occurred_at"] else "",
            t["kind"], float(t["amount"]), t["currency"], float(t["amount_uzs"]),
            t["wallet"], t["fund"] or "", t["category"] or "", t["description"] or "",
        ])

    ws2 = wb.create_sheet("P&L")
    ws2.append(["Oy", "Daromad (UZS)", "Xarajat (UZS)", "Sof (UZS)"])
    for cell in ws2[1]:
        cell.font = bold
    for r in data["pnl"]:
        inc, exp = float(r["income"]), float(r["expense"])
        ws2.append([r["month"], inc, exp, inc - exp])

    ws3 = wb.create_sheet("Balanslar")
    ws3.append(["Kassa", "Valyuta", "Balans (UZS)"])
    for cell in ws3[1]:
        cell.font = bold
    for w in data["wallets"]:
        ws3.append([w["name"], w["currency"], float(w["balance_uzs"])])
    ws3.append([])
    ws3.append(["Fond", "", "Balans (UZS)"])
    for f in data["funds"]:
        ws3.append([f["name"], "", float(f["balance_uzs"])])

    ws4 = wb.create_sheet("Qarzlar")
    ws4.append(["Yo'nalish", "Kontragent", "Valyuta", "Qoldiq", "Muddat", "Holat"])
    for cell in ws4[1]:
        cell.font = bold
    for d in data["debts"]:
        ws4.append([
            d["direction"], d["counterparty_name"], d["currency"], float(d["remaining"]),
            d["due_date"].strftime("%Y-%m-%d") if d["due_date"] else "", d["status"],
        ])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------- PDF ---------------------------
def _pdf_html(data: dict) -> str:
    today = date.today().strftime("%d.%m.%Y")

    def rows(items, render):
        return "".join(render(i) for i in items) or "<tr><td colspan='9'>—</td></tr>"

    pnl_rows = rows(
        data["pnl"],
        lambda r: f"<tr><td>{r['month']}</td>"
                  f"<td class='r'>{fmt_money(r['income'], 'UZS')}</td>"
                  f"<td class='r'>{fmt_money(r['expense'], 'UZS')}</td>"
                  f"<td class='r'>{fmt_money(Decimal(r['income']) - Decimal(r['expense']), 'UZS')}</td></tr>",
    )
    wallet_rows = rows(
        data["wallets"],
        lambda w: f"<tr><td>{w['name']} ({w['currency']})</td>"
                  f"<td class='r'>{fmt_money(w['balance_uzs'], 'UZS')}</td></tr>",
    )
    fund_rows = rows(
        data["funds"],
        lambda f: f"<tr><td>{f['name']}</td><td class='r'>{fmt_money(f['balance_uzs'], 'UZS')}</td></tr>",
    )
    cat_rows = rows(
        data["by_category"],
        lambda c: f"<tr><td>{c['name']}</td><td class='r'>{fmt_money(c['total'], 'UZS')}</td></tr>",
    )
    debt_rows = rows(
        data["debts"],
        lambda d: f"<tr><td>{d['direction']}</td><td>{d['counterparty_name']}</td>"
                  f"<td class='r'>{fmt_money(d['remaining'], d['currency'])}</td>"
                  f"<td>{d['status']}</td></tr>",
    )

    return f"""<!doctype html><html><head><meta charset="utf-8"><style>
      body {{ font-family: DejaVu Sans, sans-serif; font-size: 12px; color: #222; }}
      h1 {{ font-size: 18px; }} h2 {{ font-size: 14px; margin-top: 18px; }}
      table {{ width: 100%; border-collapse: collapse; margin-top: 6px; }}
      th, td {{ border: 1px solid #ccc; padding: 4px 6px; text-align: left; }}
      th {{ background: #f0f0f0; }} .r {{ text-align: right; }}
    </style></head><body>
      <h1>Shaxsiy moliya hisoboti</h1><p>Sana: {today}</p>
      <h2>Oylik P&amp;L</h2>
      <table><tr><th>Oy</th><th>Daromad</th><th>Xarajat</th><th>Sof</th></tr>{pnl_rows}</table>
      <h2>Kassa balanslari</h2>
      <table><tr><th>Kassa</th><th>Balans</th></tr>{wallet_rows}</table>
      <h2>Fond balanslari</h2>
      <table><tr><th>Fond</th><th>Balans</th></tr>{fund_rows}</table>
      <h2>Xarajat kategoriyalari</h2>
      <table><tr><th>Kategoriya</th><th>Summa</th></tr>{cat_rows}</table>
      <h2>Faol qarzlar</h2>
      <table><tr><th>Yo'nalish</th><th>Kontragent</th><th>Qoldiq</th><th>Holat</th></tr>{debt_rows}</table>
    </body></html>"""


async def build_pdf() -> bytes:
    """PDF hisobotini (bytes) yasaydi. weasyprint mavjud bo'lmasa RuntimeError."""
    try:
        from weasyprint import HTML
    except Exception as exc:  # tizim bog'liqliklari yo'q bo'lishi mumkin
        log.warning("weasyprint mavjud emas: %s", exc)
        raise RuntimeError("PDF eksport hozircha mavjud emas (weasyprint o'rnatilmagan).")
    data = await _collect()
    return HTML(string=_pdf_html(data)).write_pdf()
