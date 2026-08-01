import io

from openpyxl import load_workbook
from pypdf import PdfReader

from app.db import pool
from app.handlers import report
from tests.helpers import FakeCallback, FakeMessage, make_user, wallet_id


async def _seed_two_transactions(admin_id: int, wid: int):
    async with pool().acquire() as conn:
        kitob_id = await conn.fetchval("SELECT id FROM income_sources WHERE name = 'Kitob'")
        ijara_id = await conn.fetchval("SELECT id FROM expense_categories WHERE name = 'Ijara'")
        await conn.execute(
            """INSERT INTO transactions (kind, wallet_id, amount, currency, fx_rate, amount_uzs, source_id, description, created_by)
               VALUES ('income', $1, 300000, 'UZS', 1, 300000, $2, 'Kitob sotildi', $3)""",
            wid, kitob_id, admin_id,
        )
        await conn.execute(
            """INSERT INTO transactions (kind, wallet_id, amount, currency, fx_rate, amount_uzs, category_id, description, created_by)
               VALUES ('expense', $1, 100000, 'UZS', 1, 100000, $2, 'Oylik ijara', $3)""",
            wid, ijara_id, admin_id,
        )


async def test_report_requires_registration():
    m = FakeMessage(4001)
    await report.cmd_report(m)
    assert any("ro'yxatdan o'ting" in a for a in m.answers)


async def test_report_shows_period_picker():
    await make_user("viewer", 4002)
    m = FakeMessage(4002)
    await report.cmd_report(m)
    assert any("davrni tanlang" in a for a in m.answers)


async def test_period_choice_shows_format_picker():
    await make_user("admin", 4003)
    cb = FakeCallback("rep:all", 4003)
    await report.cb_report_period(cb)
    assert any("qaysi formatda" in e for e in cb.message.edits)
    assert not cb.message.documents


async def test_report_xlsx_contains_transactions_and_balances():
    admin_id = await make_user("admin", 4004)
    wid = await wallet_id()
    await _seed_two_transactions(admin_id, wid)

    cb = FakeCallback("repfmt:xlsx:all", 4004)
    await report.cb_report_format(cb)

    assert cb.message.documents, "hujjat yuborilishi kerak edi"
    doc, caption = cb.message.documents[0]
    assert "2 ta tranzaksiya" in caption

    wb = load_workbook(io.BytesIO(doc.data))
    ws = wb["Tranzaksiyalar"]
    header = [c.value for c in ws[1]]
    assert header[1] == "Turi"
    assert header[5] == "Summa (UZS)"

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    assert any(r[1] == "Kirim" and r[6] == "Kitob" for r in rows)
    assert any(r[1] == "Chiqim" and r[6] == "Ijara" for r in rows)

    balances = wb["Balanslar"]
    naqd_row = next(r for r in balances.iter_rows(values_only=True) if r[0] == "Naqd UZS")
    assert naqd_row[2] == 200000.0


async def test_report_pdf_contains_transactions_and_balances():
    admin_id = await make_user("admin", 4005)
    wid = await wallet_id()
    await _seed_two_transactions(admin_id, wid)

    cb = FakeCallback("repfmt:pdf:all", 4005)
    await report.cb_report_format(cb)

    assert cb.message.documents, "hujjat yuborilishi kerak edi"
    doc, caption = cb.message.documents[0]
    assert "2 ta tranzaksiya" in caption
    assert doc.filename.endswith(".pdf")

    reader = PdfReader(io.BytesIO(doc.data))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "Kirim" in text
    assert "Chiqim" in text
    assert "Kitob" in text
    assert "Ijara" in text
    assert "Naqd UZS" in text
    assert "200 000" in text  # kassa balansi PDF ichida to'g'ri hisoblangan


async def test_report_invalid_period_shows_alert():
    await make_user("admin", 4006)
    cb = FakeCallback("rep:noexist", 4006)
    await report.cb_report_period(cb)
    assert cb.alerts and "Noto'g'ri davr" in cb.alerts[0][0]


async def test_report_invalid_format_shows_alert():
    await make_user("admin", 4007)
    cb = FakeCallback("repfmt:docx:all", 4007)
    await report.cb_report_format(cb)
    assert cb.alerts and "Noto'g'ri so'rov" in cb.alerts[0][0]
